"""
Planilha viva da carteira — gerada todo dia pelo workflow do digest e
enviada como ARQUIVO pro seu Telegram privado. Nunca é commitada: seus
valores reais vivem no GitHub Secret CARTEIRA_PRIVADA (criptografado), não
no repositório, que é público.

O que a planilha tem:
  - Renda Fixa: cada título com a taxa contratada, acumulando dia a dia.
    As células de "dias úteis" e "valor atual" são FÓRMULAS com HOJE() —
    ou seja, mesmo um arquivo baixado ontem continua acruando o prefixado
    quando você abre hoje. Os fatores de índice (IPCA acumulado, CDI
    acumulado) são buscados das séries oficiais do Banco Central a cada
    geração.
  - Renda Variável: cotação do dia (brapi, com fallback Yahoo), variação,
    valor da posição e resultado sobre o preço médio.
  - Proventos: último provento anunciado por cota × suas cotas = quanto cai
    pra você, mais o dividend yield de 12 meses.
  - Indicadores: IPCA, CDI, Selic e dólar oficiais do dia (com a série do
    BCB anotada em comentário em cada célula).
  - Resumo: total da carteira e alocação por classe — tudo fórmula, nada
    chumbado.

Convenção de cores (padrão de modelo financeiro):
  - AZUL: dado de entrada buscado hoje (APIs) ou vindo do seu secret
  - PRETO: fórmula que recalcula sozinha
  - VERDE: referência a outra aba

Variáveis de ambiente:
  CARTEIRA_PRIVADA     - JSON com suas posições reais (ver carteira_privada.exemplo.json)
  BRAPI_TOKEN          - cotações da B3
  TELEGRAM_BOT_TOKEN   - envio do arquivo
  TELEGRAM_CHAT_ID     - chat privado que recebe
Para teste local, um arquivo carteira_privada.json (no .gitignore!) substitui o secret.
"""

import io
import json
import os
from datetime import date, datetime, timedelta

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()

BRAPI_TOKEN = os.environ.get("BRAPI_TOKEN", "")
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "")
TELEGRAM_CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID", "")

ARQUIVO_SAIDA = "carteira.xlsx"

FONTE = "Arial"
AZUL_INPUT = Font(name=FONTE, color="0000FF")          # dado de entrada
PRETO_FORMULA = Font(name=FONTE, color="000000")       # fórmula
VERDE_LINK = Font(name=FONTE, color="008000")          # referência a outra aba
TITULO = Font(name=FONTE, bold=True, size=13)
CABECALHO = Font(name=FONTE, bold=True, color="FFFFFF")
FILL_CABECALHO = PatternFill("solid", fgColor="1F4E79")

FMT_REAIS = '"R$" #,##0.00;("R$" #,##0.00)'
FMT_PCT = "0.00%"
FMT_DATA = "DD/MM/YYYY"
FMT_FATOR = "0.000000"

# Séries do SGS (Banco Central) usadas — públicas, sem chave
SGS_IPCA_MENSAL = 433   # IPCA, variação % mensal
SGS_CDI_DIARIO = 12     # CDI, % ao dia
SGS_SELIC_META = 432    # Meta Selic, % a.a.
SGS_PTAX_VENDA = 1      # Dólar PTAX venda


# ------------------------------------------------------------ Dados privados

def carrega_carteira_privada() -> dict:
    bruto = os.environ.get("CARTEIRA_PRIVADA", "").strip()
    if not bruto:
        try:
            with open("carteira_privada.json", encoding="utf-8") as f:
                bruto = f.read()
            print("Usando carteira_privada.json local (modo teste).")
        except FileNotFoundError:
            raise SystemExit(
                "CARTEIRA_PRIVADA não configurado. Copie o modelo de "
                "carteira_privada.exemplo.json, preencha com seus dados e cole "
                "num GitHub Secret com esse nome (nunca num arquivo do repositório)."
            )
    return json.loads(bruto)


# ------------------------------------------------------- Séries do BCB (SGS)

def _sgs(codigo: int, data_inicial: date | None = None) -> list[dict]:
    params = {"formato": "json"}
    if data_inicial:
        params["dataInicial"] = data_inicial.strftime("%d/%m/%Y")
        params["dataFinal"] = date.today().strftime("%d/%m/%Y")
        url = f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{codigo}/dados"
    else:
        url = f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{codigo}/dados/ultimos/1"
    resp = requests.get(url, params=params, timeout=20)
    resp.raise_for_status()
    return resp.json()


def fator_ipca_acumulado(data_aplicacao: date) -> tuple[float, str]:
    """Produto de (1 + IPCA mensal) dos meses FECHADOS e já publicados
    posteriores ao mês da aplicação. Convenção simplificada: o mês da
    aplicação não entra (não há pro-rata); o fator anda quando o IBGE
    publica o mês seguinte."""
    dados = _sgs(SGS_IPCA_MENSAL, data_inicial=data_aplicacao)
    fator, ultimo_mes = 1.0, "-"
    for d in dados:
        mes = datetime.strptime(d["data"], "%d/%m/%Y").date()
        if (mes.year, mes.month) > (data_aplicacao.year, data_aplicacao.month):
            fator *= 1 + float(d["valor"]) / 100
            ultimo_mes = mes.strftime("%m/%Y")
    return fator, ultimo_mes


def fator_cdi_acumulado(data_aplicacao: date, percentual: float) -> tuple[float, str]:
    """Produto de (1 + CDI diário × percentual) dos dias úteis desde a
    aplicação (exclusive) até o último dado publicado."""
    dados = _sgs(SGS_CDI_DIARIO, data_inicial=data_aplicacao + timedelta(days=1))
    fator, ultimo_dia = 1.0, "-"
    for d in dados:
        fator *= 1 + (float(d["valor"]) / 100) * (percentual / 100)
        ultimo_dia = d["data"]
    return fator, ultimo_dia


def indicadores_do_dia() -> dict:
    """Último valor de cada indicador; item que falhar vira None (a aba
    Indicadores mostra 'indisponível' em vez de quebrar a planilha)."""
    resultado = {}
    for chave, codigo in [("selic", SGS_SELIC_META), ("ptax", SGS_PTAX_VENDA),
                          ("ipca_mes", SGS_IPCA_MENSAL), ("cdi_dia", SGS_CDI_DIARIO)]:
        try:
            d = _sgs(codigo)[-1]
            resultado[chave] = (float(d["valor"]), d["data"])
        except Exception as exc:
            print(f"Indicador {chave} (SGS {codigo}) indisponível: {exc}")
            resultado[chave] = None
    return resultado


# ------------------------------------------------------------------ Cotações

def cotacao(ticker: str) -> dict | None:
    """Preço e variação do dia. brapi primeiro; Yahoo como fallback."""
    if BRAPI_TOKEN:
        try:
            resp = requests.get(
                f"https://brapi.dev/api/quote/{ticker}",
                headers={"Authorization": f"Bearer {BRAPI_TOKEN}"},
                timeout=20,
            )
            resp.raise_for_status()
            r = resp.json()["results"][0]
            return {"preco": r.get("regularMarketPrice"),
                    "variacao": r.get("regularMarketChangePercent"),
                    "fonte": "brapi.dev"}
        except Exception as exc:
            print(f"brapi falhou pra {ticker} ({exc}), tentando Yahoo...")
    try:
        import yfinance as yf

        hist = yf.Ticker(f"{ticker}.SA").history(period="5d")["Close"].dropna()
        if len(hist) >= 2:
            return {"preco": float(hist.iloc[-1]),
                    "variacao": (float(hist.iloc[-1]) / float(hist.iloc[-2]) - 1) * 100,
                    "fonte": "Yahoo Finance"}
    except Exception as exc:
        print(f"Yahoo também falhou pra {ticker}: {exc}")
    return None


def proventos(ticker: str) -> dict | None:
    """Último provento por cota e soma de 12 meses, via Yahoo Finance."""
    try:
        import yfinance as yf

        divs = yf.Ticker(f"{ticker}.SA").dividends
        if divs is None or divs.empty:
            return None
        ultimo_quando = divs.index[-1]
        soma_12m = float(divs[divs.index >= ultimo_quando - timedelta(days=365)].sum())
        return {"ultimo": float(divs.iloc[-1]),
                "data": ultimo_quando.strftime("%d/%m/%Y"),
                "soma_12m": soma_12m}
    except Exception as exc:
        print(f"Proventos indisponíveis pra {ticker}: {exc}")
        return None


# ---------------------------------------------------------------- Montagem

def _cabecalho(ws, linha: int, colunas: list[str]):
    for i, nome in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.font = CABECALHO
        c.fill = FILL_CABECALHO
        c.alignment = Alignment(horizontal="center")


def _largura(ws, larguras: list[int]):
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def gera_planilha(carteira: dict) -> Workbook:
    wb = Workbook()

    # ---------------------------------------------------------- Indicadores
    ws_ind = wb.active
    ws_ind.title = "Indicadores"
    ws_ind["A1"] = "Indicadores oficiais do dia"
    ws_ind["A1"].font = TITULO
    ind = indicadores_do_dia()
    linhas_ind = [
        ("Meta Selic (% a.a.)", "selic", f"Fonte: Banco Central, série SGS {SGS_SELIC_META}"),
        ("CDI (% ao dia)", "cdi_dia", f"Fonte: Banco Central, série SGS {SGS_CDI_DIARIO}"),
        ("IPCA do último mês (%)", "ipca_mes", f"Fonte: Banco Central/IBGE, série SGS {SGS_IPCA_MENSAL}"),
        ("Dólar PTAX venda (R$)", "ptax", f"Fonte: Banco Central, série SGS {SGS_PTAX_VENDA}"),
    ]
    for i, (rotulo, chave, fonte) in enumerate(linhas_ind, start=3):
        ws_ind.cell(row=i, column=1, value=rotulo).font = Font(name=FONTE)
        valor = ind.get(chave)
        c = ws_ind.cell(row=i, column=2, value=valor[0] if valor else "indisponível")
        c.font = AZUL_INPUT
        c.comment = Comment(
            f"{fonte}. Dado de {valor[1] if valor else '—'}. Buscado automaticamente a cada geração.",
            "carteira_excel.py",
        )
        if valor:
            c.number_format = "#,##0.0000"
    ws_ind["A8"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} (UTC do runner)"
    ws_ind["A8"].font = Font(name=FONTE, italic=True, size=9)
    ws_ind["A10"] = "Legenda: células AZUIS = dado buscado hoje (APIs/secret) · PRETAS = fórmula que recalcula sozinha · VERDES = referência a outra aba"
    ws_ind["A10"].font = Font(name=FONTE, italic=True, size=9)
    _largura(ws_ind, [34, 16])

    # ----------------------------------------------------------- Renda Fixa
    ws_rf = wb.create_sheet("Renda Fixa")
    ws_rf["A1"] = "Renda fixa — acumulando dia a dia"
    ws_rf["A1"].font = TITULO
    cols = ["Nome", "Indexador", "Taxa contratada (a.a.)", "Data aplicação", "Vencimento",
            "Valor aplicado", "Fator do índice", "Dias úteis", "Valor atual", "Rendimento", "% no período"]
    _cabecalho(ws_rf, 3, cols)
    linha = 4
    for pos in carteira.get("renda_fixa", []):
        aplic = datetime.strptime(pos["data_aplicacao"], "%Y-%m-%d").date()
        indexador = pos["indexador"].lower()

        ws_rf.cell(row=linha, column=1, value=pos["nome"]).font = AZUL_INPUT
        ws_rf.cell(row=linha, column=2, value=indexador).font = AZUL_INPUT

        taxa = pos.get("taxa_aa")
        c_taxa = ws_rf.cell(row=linha, column=3)
        if indexador == "cdi":
            c_taxa.value = pos.get("percentual_cdi", 100) / 100
            c_taxa.number_format = "0%"
            c_taxa.comment = Comment("Percentual do CDI contratado (do seu secret).", "carteira_excel.py")
        else:
            c_taxa.value = (taxa or 0) / 100
            c_taxa.number_format = FMT_PCT
            c_taxa.comment = Comment(
                "Taxa anual contratada (do seu secret). No IPCA+, é o spread ACIMA da inflação.",
                "carteira_excel.py",
            )
        c_taxa.font = AZUL_INPUT

        c_aplic = ws_rf.cell(row=linha, column=4, value=aplic)
        c_aplic.number_format = FMT_DATA
        c_aplic.font = AZUL_INPUT
        if pos.get("vencimento"):
            c_venc = ws_rf.cell(row=linha, column=5,
                                value=datetime.strptime(pos["vencimento"], "%Y-%m-%d").date())
            c_venc.number_format = FMT_DATA
            c_venc.font = AZUL_INPUT

        c_val = ws_rf.cell(row=linha, column=6, value=float(pos["valor_aplicado"]))
        c_val.number_format = FMT_REAIS
        c_val.font = AZUL_INPUT

        # fator do índice (IPCA ou CDI acumulado) — buscado do BCB a cada geração
        c_fator = ws_rf.cell(row=linha, column=7)
        c_fator.font = AZUL_INPUT
        c_fator.number_format = FMT_FATOR
        if indexador == "ipca+":
            fator, ate = fator_ipca_acumulado(aplic)
            c_fator.value = fator
            c_fator.comment = Comment(
                f"IPCA acumulado desde a aplicação (meses fechados até {ate}), série SGS "
                f"{SGS_IPCA_MENSAL} do BCB. O mês corrente entra quando o IBGE publica.",
                "carteira_excel.py",
            )
        elif indexador == "cdi":
            fator, ate = fator_cdi_acumulado(aplic, pos.get("percentual_cdi", 100))
            c_fator.value = fator
            c_fator.comment = Comment(
                f"CDI acumulado × percentual contratado, dias úteis até {ate}, série SGS "
                f"{SGS_CDI_DIARIO} do BCB. Atualiza a cada geração da planilha.",
                "carteira_excel.py",
            )
        else:
            c_fator.value = 1.0
            c_fator.comment = Comment("Prefixado não usa índice — fator fixo em 1.", "carteira_excel.py")

        # dias úteis e valor atual: FÓRMULAS com HOJE() — o arquivo continua
        # acruando o componente prefixado mesmo dias depois de gerado
        c_du = ws_rf.cell(row=linha, column=8, value=f"=NETWORKDAYS(D{linha},TODAY())")
        c_du.font = PRETO_FORMULA
        c_du.comment = Comment(
            "Dias úteis (seg-sex) da aplicação até hoje. Aproximação: não desconta "
            "feriados — pequena diferença frente à convenção DU/252 oficial.",
            "carteira_excel.py",
        )

        if indexador == "cdi":
            formula_valor = f"=F{linha}*G{linha}"
        else:  # prefixado e ipca+ capitalizam a taxa anual por DU/252
            formula_valor = f"=F{linha}*G{linha}*(1+C{linha})^(H{linha}/252)"
        c_atual = ws_rf.cell(row=linha, column=9, value=formula_valor)
        c_atual.number_format = FMT_REAIS
        c_atual.font = PRETO_FORMULA

        c_rend = ws_rf.cell(row=linha, column=10, value=f"=I{linha}-F{linha}")
        c_rend.number_format = FMT_REAIS
        c_rend.font = PRETO_FORMULA
        c_pct = ws_rf.cell(row=linha, column=11, value=f"=IF(F{linha}=0,0,I{linha}/F{linha}-1)")
        c_pct.number_format = FMT_PCT
        c_pct.font = PRETO_FORMULA
        linha += 1

    ultima_rf = linha - 1
    if ultima_rf >= 4:
        ws_rf.cell(row=linha + 1, column=1, value="Total").font = Font(name=FONTE, bold=True)
        for col in (6, 9, 10):
            letra = get_column_letter(col)
            c = ws_rf.cell(row=linha + 1, column=col, value=f"=SUM({letra}4:{letra}{ultima_rf})")
            c.number_format = FMT_REAIS
            c.font = Font(name=FONTE, bold=True)
    _largura(ws_rf, [34, 11, 12, 13, 13, 14, 12, 10, 14, 13, 12])

    # ------------------------------------------------------- Renda Variável
    ws_rv = wb.create_sheet("Renda Variável")
    ws_rv["A1"] = "Renda variável — cotações de hoje"
    ws_rv["A1"].font = TITULO
    _cabecalho(ws_rv, 3, ["Ticker", "Quantidade", "Preço médio", "Cotação", "Variação (dia)",
                          "Valor da posição", "Resultado"])
    linha = 4
    for pos in carteira.get("renda_variavel", []):
        ticker = pos["ticker"]
        q = cotacao(ticker)
        ws_rv.cell(row=linha, column=1, value=ticker).font = AZUL_INPUT
        ws_rv.cell(row=linha, column=2, value=pos.get("quantidade", 0)).font = AZUL_INPUT
        if pos.get("preco_medio") is not None:
            c_pm = ws_rv.cell(row=linha, column=3, value=float(pos["preco_medio"]))
            c_pm.number_format = FMT_REAIS
            c_pm.font = AZUL_INPUT
        c_cot = ws_rv.cell(row=linha, column=4, value=q["preco"] if q else "indisponível")
        c_cot.font = AZUL_INPUT
        if q:
            c_cot.number_format = FMT_REAIS
            c_cot.comment = Comment(f"Fonte: {q['fonte']}, no momento da geração.", "carteira_excel.py")
            c_var = ws_rv.cell(row=linha, column=5, value=(q["variacao"] or 0) / 100)
            c_var.number_format = FMT_PCT
            c_var.font = AZUL_INPUT
        c_pos = ws_rv.cell(row=linha, column=6, value=f"=B{linha}*D{linha}")
        c_pos.number_format = FMT_REAIS
        c_pos.font = PRETO_FORMULA
        c_res = ws_rv.cell(
            row=linha, column=7,
            value=f'=IF(C{linha}="",0,(D{linha}-C{linha})*B{linha})',
        )
        c_res.number_format = FMT_REAIS
        c_res.font = PRETO_FORMULA
        linha += 1
    ultima_rv = linha - 1
    if ultima_rv >= 4:
        ws_rv.cell(row=linha + 1, column=1, value="Total").font = Font(name=FONTE, bold=True)
        for col in (6, 7):
            letra = get_column_letter(col)
            c = ws_rv.cell(row=linha + 1, column=col, value=f"=SUM({letra}4:{letra}{ultima_rv})")
            c.number_format = FMT_REAIS
            c.font = Font(name=FONTE, bold=True)
    _largura(ws_rv, [10, 12, 13, 13, 13, 16, 14])

    # ------------------------------------------------------------ Proventos
    ws_pv = wb.create_sheet("Proventos")
    ws_pv["A1"] = "Proventos — quanto cai pra você"
    ws_pv["A1"].font = TITULO
    _cabecalho(ws_pv, 3, ["Ticker", "Quantidade", "Último provento/cota", "Anunciado em",
                          "Você recebe", "Soma 12m/cota", "DY 12m"])
    linha = 4
    for pos in carteira.get("renda_variavel", []):
        ticker = pos["ticker"]
        p = proventos(ticker)
        ws_pv.cell(row=linha, column=1, value=ticker).font = AZUL_INPUT
        c_q = ws_pv.cell(row=linha, column=2, value=f"='Renda Variável'!B{linha}")
        c_q.font = VERDE_LINK
        if p:
            c_u = ws_pv.cell(row=linha, column=3, value=p["ultimo"])
            c_u.number_format = FMT_REAIS
            c_u.font = AZUL_INPUT
            c_u.comment = Comment(
                "Último provento ANUNCIADO por cota (Yahoo Finance). Proventos futuros não "
                "existem em fonte pública confiável — o que dá pra saber é o anunciado e o histórico.",
                "carteira_excel.py",
            )
            ws_pv.cell(row=linha, column=4, value=p["data"]).font = AZUL_INPUT
            c_rec = ws_pv.cell(row=linha, column=5, value=f"=B{linha}*C{linha}")
            c_rec.number_format = FMT_REAIS
            c_rec.font = PRETO_FORMULA
            c_12 = ws_pv.cell(row=linha, column=6, value=p["soma_12m"])
            c_12.number_format = FMT_REAIS
            c_12.font = AZUL_INPUT
            c_dy = ws_pv.cell(
                row=linha, column=7,
                value=f"=IF('Renda Variável'!D{linha}=0,0,F{linha}/'Renda Variável'!D{linha})",
            )
            c_dy.number_format = FMT_PCT
            c_dy.font = VERDE_LINK
        else:
            ws_pv.cell(row=linha, column=3, value="indisponível").font = AZUL_INPUT
        linha += 1
    _largura(ws_pv, [10, 12, 18, 13, 13, 14, 10])

    # --------------------------------------------------------------- Resumo
    ws_res = wb.create_sheet("Resumo", 0)
    ws_res["A1"] = "Resumo da carteira"
    ws_res["A1"].font = TITULO
    ws_res["A2"] = f"Atualizado em {datetime.now().strftime('%d/%m/%Y')} — renda fixa continua acruando ao abrir (fórmulas com HOJE())"
    ws_res["A2"].font = Font(name=FONTE, italic=True, size=9)

    _cabecalho(ws_res, 4, ["Classe", "Valor atual", "% da carteira"])
    ws_res["A5"] = "Renda fixa"
    ws_res["A5"].font = Font(name=FONTE)
    c = ws_res.cell(row=5, column=2, value=f"=SUM('Renda Fixa'!I4:I{max(ultima_rf, 4)})")
    c.number_format = FMT_REAIS
    c.font = VERDE_LINK
    ws_res["A6"] = "Renda variável"
    ws_res["A6"].font = Font(name=FONTE)
    c = ws_res.cell(row=6, column=2, value=f"=SUM('Renda Variável'!F4:F{max(ultima_rv, 4)})")
    c.number_format = FMT_REAIS
    c.font = VERDE_LINK
    ws_res["A7"] = "Total"
    ws_res["A7"].font = Font(name=FONTE, bold=True)
    c = ws_res.cell(row=7, column=2, value="=B5+B6")
    c.number_format = FMT_REAIS
    c.font = Font(name=FONTE, bold=True)
    for r in (5, 6):
        c = ws_res.cell(row=r, column=3, value=f"=IF($B$7=0,0,B{r}/$B$7)")
        c.number_format = FMT_PCT
        c.font = PRETO_FORMULA
    _largura(ws_res, [20, 16, 14])

    return wb


# ------------------------------------------------------------------- Envio

def envia_telegram(caminho: str) -> None:
    if not (TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID):
        print("Sem TELEGRAM_BOT_TOKEN/TELEGRAM_CHAT_ID — arquivo gerado mas não enviado (modo teste).")
        return
    with open(caminho, "rb") as f:
        resp = requests.post(
            f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument",
            data={
                "chat_id": TELEGRAM_CHAT_ID,
                "caption": f"📒 Sua carteira de {datetime.now().strftime('%d/%m')} — renda fixa acruada, cotações e proventos do dia.",
            },
            files={"document": (f"carteira_{datetime.now().strftime('%Y-%m-%d')}.xlsx", f)},
            timeout=60,
        )
    if resp.ok:
        print("carteira.xlsx enviada no Telegram.")
    else:
        print("Falha enviando a planilha no Telegram:", resp.text)


def main() -> None:
    carteira = carrega_carteira_privada()
    wb = gera_planilha(carteira)
    wb.save(ARQUIVO_SAIDA)
    print(f"{ARQUIVO_SAIDA} gerado.")
    envia_telegram(ARQUIVO_SAIDA)


if __name__ == "__main__":
    main()
